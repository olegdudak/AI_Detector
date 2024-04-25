import numpy as np
import tkinter as tk
import pandas as pd
import docx
import re
from sklearn.feature_extraction.text import CountVectorizer
from sklearn.model_selection import train_test_split
from sklearn.naive_bayes import MultinomialNB
from sklearn.metrics import accuracy_score, f1_score, confusion_matrix
from tkinter import Label, Text, Button, messagebox
from tkinter import filedialog
from PyPDF2 import PdfReader


class TextDetectorGUI:
    def __init__(self, master, model, vectorizer, base_training_data):
        self.master = master
        master.title("AI Detector")

        self.label = Label(master, text="Add text: ")
        self.label.pack(pady=(10, 10))

        self.text_frame = tk.Frame(master)
        self.text_frame.pack()

        self.text_entry = Text(self.text_frame, height=20, width=100)
        self.text_entry.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self.scrollbar = tk.Scrollbar(self.text_frame, command=self.text_entry.yview)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.text_entry.config(yscrollcommand=self.scrollbar.set)

        self.analyse_button = Button(master, text="Scan", command=self.analyse_text)
        self.analyse_button.pack(side=tk.LEFT, padx=(10, 10))

        self.result_label = Label(master, text="")
        self.result_label.pack(pady=(15, 15))

        checkboxes_frame = tk.Frame(master)
        checkboxes_frame.pack(side=tk.RIGHT, anchor=tk.CENTER)

        self.reference_var = tk.IntVar()
        self.reference_checkbox = tk.Checkbutton(checkboxes_frame, text="References", variable=self.reference_var, command=self.update_reference)
        self.reference_checkbox.pack(pady=(0, 0))

        self.sources_var = tk.IntVar()
        self.sources_checkbox = tk.Checkbutton(checkboxes_frame, text="Sources", variable=self.sources_var, command=self.update_source)
        self.sources_checkbox.pack(pady=(0, 10), padx=(0, 17))

        self.upload_button = Button(checkboxes_frame, text="Upload Document", command=self.upload_document)
        self.upload_button.pack(side=tk.LEFT, padx=(10, 10), pady=(0, 10))

        self.model = model
        self.vectorizer = vectorizer
        self.base_training_data = base_training_data
        self.highlight_colour = "yellow"

    def upload_document(self):
        file_path = filedialog.askopenfilename(filetypes=[("Text files", "*.txt;*.docx;*.pdf")])
        if file_path:
            with open(file_path, "rb") as file:
                if file_path.lower().endswith(".txt"):
                    content = file.read().decode("utf-8")
                elif file_path.lower().endswith(".docx"):
                    doc = docx.Document(file)
                    content = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                elif file_path.lower().endswith(".pdf"):
                    pdf_reader = PdfReader(file)
                    content = ""
                    for page_num in range(len(pdf_reader.pages)):
                        page = pdf_reader.pages[page_num]
                        content += page.extract_text()
                else:
                    messagebox.showwarning("Invalid File", "Upload a valid text, Word, or PDF document...")
                    return

            self.text_entry.delete("1.0", tk.END)
            self.text_entry.insert(tk.END, content)


    def update_reference(self):
        if self.reference_var.get() == 1:
            return 0
        else:
            return 1

    def update_source(self):
        if self.sources_var.get() == 1:
            return 0
        else:
            return 1

    def calculate_score(self, percentage):
        if percentage <= 40:
            return "Perfect"
        elif 40 < percentage <= 60:
            return "May contain fragments of AI text"
        else:
            return "Bad"

    def analyse_text(self):
        input_text = self.text_entry.get("1.0", tk.END).strip()

        reference_value = self.update_reference()
        source_value = self.update_source()

        input_vector_text = self.vectorizer.transform([input_text]).toarray()
        reference_source_values = np.array([[reference_value, source_value]])
        input_vector = np.concatenate((input_vector_text, reference_source_values), axis=1)

        ai_words = 0
        human_words = 0

        sentences = re.split(r"(?<!\w\.\w.)(?<![A-Z][a-z]\.)(?<=\!|\?|\.|\:)", input_text)

        self.text_entry.insert(tk.END, "\n\nSuspicious Text:\n")
        for sentence in sentences:
            sentence_vector = self.vectorizer.transform([sentence]).toarray()
            sentence_vector = np.concatenate((sentence_vector, reference_source_values), axis=1)
            sentence_probability = self.model.predict_proba(sentence_vector)[0][1] * 100

            if sentence_probability > 50:
                ai_words += len(sentence.split())
                self.text_entry.insert(tk.END, sentence, 'highlighted')
            else:
                human_words += len(sentence.split())
                self.text_entry.insert(tk.END, sentence)
            self.text_entry.insert(tk.END, '\n')
        self.text_entry.tag_configure('highlighted', background='yellow')

        total_words = ai_words + human_words
        ai_percentage = (ai_words / total_words) * 100 if total_words != 0 else 0

        ai_result = round(ai_percentage, 1)
        score = self.calculate_score(ai_result)
        result_text = f"AI-text : {ai_result}%\nRating: {score}"
        self.result_label.config(text=result_text)


def load_training_data(file_path):
    df = pd.read_csv(file_path, encoding="latin-1")
    return df


def train_model(training_data):
    vectorizer = CountVectorizer()

    X_text = vectorizer.fit_transform(training_data["TEXT"]).toarray()
    X_numeric = np.array(training_data[["REFERENCE", "SOURCE"]], dtype=float)
    X = np.concatenate((X_text, X_numeric), axis=1)

    y = np.array(training_data["ID"])

    X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=0)
    trained_model = MultinomialNB()
    trained_model.fit(X_train, y_train)

    y_prediction = trained_model.predict(X_test)
    acc = accuracy_score(y_test, y_prediction)
    f1 = f1_score(y_test, y_prediction)
    confusion = confusion_matrix(y_test, y_prediction)

    print(f"Accuracy: {acc}")
    print(f"F1: {f1}")
    print(f"Confusion Matrix: ")
    print(confusion)

    return trained_model, vectorizer


if __name__ == "__main__":
    training_data = load_training_data("TextBigData.csv")
    model, vectorizer = train_model(training_data)

    root = tk.Tk()
    app = TextDetectorGUI(root, model, vectorizer, training_data)
    root.mainloop()
